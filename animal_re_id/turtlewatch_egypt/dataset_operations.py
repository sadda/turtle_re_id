import requests
import pandas as pd
import os
from docx import Document
from docx.shared import Pt
import pandas as pd
from tqdm import tqdm
from wildlife_datasets.datasets import TurtlewatchEgypt_New, utils


IMAGE_EXTENSIONS = (
    ".jpg", ".jpeg",
    ".png",
    ".bmp",
    ".tif", ".tiff",
    ".webp",
    ".gif",
    ".heic", ".heif",
    ".avif"
)

def merge_datasets(
        master_root,
        roots,
        load_segmentation=True,
        encounter_prefixes=None,
        **kwargs
        ):

    dfs = []
    if encounter_prefixes is None:
        encounter_prefixes = range(len(roots))
    for root, encounter_prefix in zip(roots, encounter_prefixes):
        dataset = TurtlewatchEgypt_New(f'{master_root}/{root}', load_segmentation=load_segmentation)
        dataset.df['path'] = root.split('/')[-1] + '/' + dataset.df['path']
        dataset.df['encounter_id'] = dataset.df['encounter_id'].apply(lambda x: f'{encounter_prefix}_{x}')
        dfs.append(dataset.df)
    df = pd.concat(dfs).reset_index(drop=True)
    # TODO: image_id should be unique
    df['image_id'] = range(len(df))
    return TurtlewatchEgypt_New(root=master_root, df=df, **kwargs)

def load_citizen_data(data):
    data = pd.read_csv("citizen_files.csv")

    # Convert dates
    data["submit_time"] = pd.to_datetime(data["submit_time"])
    data["date-227"] = pd.to_datetime(data["date-227"])

    # Merge multiple date options
    data["date"] = data["submit_time"].combine_first(data["date-227"])
    data["year"] = data["date"].dt.year
    data["month"] = data["date"].dt.month
    data["day"] = data["date"].dt.day

    # Merge multiple author and emails options
    data["author"] = data["NOME"].combine_first(data["your-name"])
    data["email"] = data["EMAIL"].combine_first(data["email-185"])

    # Fill nans
    data["author"] = data["author"].fillna("unknown")
    data["email"] = data["email"].fillna("unknown")

    # Get folder and sightings
    data["folder"] = [get_folder(d) for _, d in data.iterrows()]
    for _, data_folder in data.groupby(["folder"]):
        data.loc[data_folder.index, "sighting"] = list(range(1, len(data_folder)+1))
    data["sighting"] = data["sighting"].astype(int)

    return data

def download_dataset(root, data):
    metadata = pd.DataFrame()
    for encounter, (_, d) in tqdm(enumerate(data.iterrows()), total=len(data)):
        urls = d["upload-file-731"].split("\n")
        folder = d["folder"]
        sighting = d["sighting"]
        folder_full = f"{root}/{folder}/SIGHTING #{sighting}"

        save_paths = download_files(urls, folder_full)
        save_paths = [os.path.relpath(p, root) for p in save_paths]
        
        create_info(d, folder_full)

        metadata_part = {
            "path": save_paths,
            "identity": "unknown",
            "encounter": encounter,
        }
        metadata = pd.concat((metadata, pd.DataFrame(metadata_part)))
    metadata = metadata.reset_index(drop=True)
    metadata["image_id"] = utils.get_persistent_id(metadata["path"])
    metadata.to_csv(os.path.join(root, "metadata.csv"), index=False)


def get_folder(d):
    year = d["year"]
    month = d["month"]
    day = d["day"]

    folder1 = f"{year}_{month:02d}_{day:02d}"
    folder2 = d["author"]
    return f"{folder1}/{folder2}"

def download_files(urls, download_folder, exts=IMAGE_EXTENSIONS):
    os.makedirs(download_folder, exist_ok=True)

    save_paths = []
    for url in urls:
        if not url.lower().endswith(exts):
            print(f"Skipping non-image url: {url}")
            continue
        file_name = url.split("/")[-1]
        save_path = os.path.join(download_folder, file_name)
        save_paths.append(save_path)
        if os.path.exists(save_path):
            continue
        try:
            response = requests.get(url, timeout=20)
            if response.status_code == 200:
                with open(save_path, "wb") as f:
                    f.write(response.content)
            else:
                print(f"Failed: {url}")
        except Exception as e:
            print(f"Error downloading {url}: {e}")
    return save_paths

def add_run_break(p, text1, text2=None):
    if not pd.isnull(text2):
        r = p.add_run(f"{text1}: {text2}")
    else:
        r = p.add_run(f"{text1}:")
    r.add_break()

def create_info(d, save_folder):
    doc = Document()

    font = doc.styles["Normal"].font
    font.name = "Arial"
    font.size = Pt(11)

    p = doc.add_paragraph()
    add_run_break(p, "REQUIRED DATA")
    add_run_break(p, "From", d["email"])
    add_run_break(p, "Photographer", d["author"])
    add_run_break(p, "Date", f'{d["year"]}-{d["month"]:02d}-{d["day"]:02d}')
    add_run_break(p, "Town", d["town-785"])
    add_run_break(p, "Location", d["location-785"])
    add_run_break(p, "")
    add_run_break(p, "OPTIONAL DATA")
    add_run_break(p, "Dive centre/independent", d["dive-785"])
    add_run_break(p, "Time", d["time-299"])
    add_run_break(p, "Depth", d["depth-364"])
    add_run_break(p, "Temperature", d["degrees-364"])
    add_run_break(p, "Activity", d["radio-602"])
    add_run_break(p, "Species", d["species-954"])
    add_run_break(p, "Size", d["size-648"])
    add_run_break(p, "Sex", d["sex-786"])
    add_run_break(p, "Comments", d["textarea-268"])

    doc.save(f"{save_folder}/info.docx")